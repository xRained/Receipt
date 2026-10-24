import csv
import io
import json
import re
import sqlite3
import uuid
import hashlib
from datetime import date, datetime
from pathlib import Path
from difflib import SequenceMatcher
from contextlib import contextmanager

import openpyxl
from matching import plan_matches, review_matches, identity
from workbooks import audit, normalized_receipts, account_entries
from flask import Flask, jsonify, render_template, request, send_file, send_from_directory
from version import version
import os

ROOT = Path(__file__).parent
HOSTED_PREVIEW = os.environ.get('VERCEL') == '1'
DATA = Path(os.environ.get('RECEIPT_DATA_DIR', '/tmp/jolyn-tae' if HOSTED_PREVIEW else str(ROOT / 'data')))
DATA.mkdir(parents=True,exist_ok=True)
(DATA / 'uploads').mkdir(exist_ok=True)
app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 30 * 1024 * 1024
app.config['HOSTED_PREVIEW']=HOSTED_PREVIEW

@app.before_request
def protect_hosted_preview():
    if HOSTED_PREVIEW and request.method not in ['GET','HEAD','OPTIONS']:
        return jsonify(error='This hosted preview is read-only. Persistent storage and access control must be configured before saving financial records.'),503
RUNNING_VERSION=version()
FIELDS = ['date', 'company', 'reference', 'method', 'tin', 'address', 'amount', 'fee', 'remarks', 'code', 'account', 'project', 'receipt_no']

@contextmanager
def db():
    c = sqlite3.connect(DATA / 'receipts.sqlite3')
    c.row_factory = sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    try:
        with c:
            yield c
    finally:
        c.close()

with db() as c:
    c.executescript('''
    CREATE TABLE IF NOT EXISTS records(id TEXT PRIMARY KEY, kind TEXT NOT NULL, body TEXT NOT NULL, source TEXT, source_row INTEGER, UNIQUE(source, source_row));
    CREATE TABLE IF NOT EXISTS allocations(payment TEXT REFERENCES records(id), receipt TEXT REFERENCES records(id), cents INTEGER CHECK(cents>0), PRIMARY KEY(payment,receipt));
    CREATE TABLE IF NOT EXISTS suppliers(name TEXT PRIMARY KEY, tin TEXT, address TEXT);
    CREATE TABLE IF NOT EXISTS accounts(code TEXT PRIMARY KEY, title TEXT);
    CREATE TABLE IF NOT EXISTS match_audit(payment TEXT, receipt TEXT, reason TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP, PRIMARY KEY(payment,receipt));
    CREATE TABLE IF NOT EXISTS match_blocks(payment TEXT, receipt TEXT, PRIMARY KEY(payment,receipt));
    CREATE TABLE IF NOT EXISTS supplier_aliases(payment_identity TEXT, receipt_identity TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP, PRIMARY KEY(payment_identity,receipt_identity));
    CREATE TABLE IF NOT EXISTS workspace_settings(key TEXT PRIMARY KEY,value TEXT);
    ''')

def cents(x):
    from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
    try:
        n = Decimal(str(x or 0).replace(',', ''))
        if not n.is_finite() or n < 0:
            raise ValueError('Amount must be a finite positive number')
        return int((n * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    except InvalidOperation:
        raise ValueError('Invalid amount')

def clean(x):
    if x is None: return ''
    if isinstance(x, (datetime, date)): return x.isoformat()[:10]
    if isinstance(x, float) and x.is_integer(): return str(int(x))
    return str(x).strip()

def records(c):
    result = []
    allocations=list(c.execute('SELECT payment,receipt,cents FROM allocations'))
    balances={}
    linked={}
    returned={r['payment']:r['total'] for r in c.execute('SELECT payment,SUM(cents) AS total FROM returns GROUP BY payment')}
    for a in allocations:
        for rid in [a['payment'],a['receipt']]:balances[rid]=balances.get(rid,0)+a['cents']
        linked.setdefault(a['payment'],[]).append(a['receipt'])
    for r in c.execute('SELECT * FROM records'):
        b = json.loads(r['body'])
        b.update(id=r['id'], kind=r['kind'], source=r['source'], source_row=r['source_row'])
        b.setdefault('treatment','advance' if b.get('collection')=='COS_Allowance' else 'no_receipt' if b.get('collection')=='Non Receipt' else 'personal' if b.get('personal') else 'business')
        b['allocated'] = balances.get(r['id'],0) / 100
        b['settlement_amount'] = b.get('settlement_amount', b.get('amount', 0))
        b['returned']=returned.get(r['id'],0)/100
        b['remaining'] = round(float(b['settlement_amount'] or 0) - b['allocated'] - b['returned'], 2)
        result.append(b)
    by_id = {r['id']:r for r in result}
    for b in result:
        b['linked_receipt_numbers'] = ', '.join(dict.fromkeys(by_id[rid].get('receipt_no') or '(no receipt number)' for rid in linked.get(b['id'],[])))
    return result

def auto_match():
    with db() as c:
        workflow=c.execute("SELECT value FROM workspace_settings WHERE key='workflow'").fetchone()
        if workflow and workflow['value'] in ['online_material_receipt_presence','materials_intake_receipt_presence']:return {'links':0,'payments':0,'amount':0,'note':'Receipt presence only; no payment allocations'}
        c.execute('BEGIN IMMEDIATE')
        blocked = [tuple(r) for r in c.execute('SELECT payment,receipt FROM match_blocks')]
        aliases = [tuple(r) for r in c.execute('SELECT payment_identity,receipt_identity FROM supplier_aliases')]
        planned = plan_matches(records(c), blocked, aliases)
        for a in planned:
            c.execute('INSERT INTO allocations VALUES(?,?,?) ON CONFLICT(payment,receipt) DO UPDATE SET cents=cents+excluded.cents', (a['payment'],a['receipt'],a['cents']))
            c.execute('INSERT OR REPLACE INTO match_audit(payment,receipt,reason) VALUES(?,?,?)', (a['payment'],a['receipt'],a['reason']))
    return {'links': len(planned), 'payments':len({a['payment'] for a in planned}), 'amount':sum(a['cents'] for a in planned)/100}

@app.post('/api/auto-match')
def run_matching():
    return jsonify(auto_match())

def enrich_workbook(workbook):
    """One-time layout repair; preserve subsequent user edits and allocations."""
    updated=0
    with db() as c:
        for code,title in account_entries(workbook):
            c.execute('INSERT OR IGNORE INTO accounts VALUES(?,?)',(code,title))
        for data in normalized_receipts(workbook):
            row=c.execute('SELECT * FROM records WHERE source=? AND source_row=?',(data['source'],data['source_row'])).fetchone()
            if not row:continue
            body=json.loads(row['body'])
            if 'receipt_numbers' in body:continue
            allocated=c.execute('SELECT COALESCE(SUM(cents),0) FROM allocations WHERE receipt=?',(row['id'],)).fetchone()[0]
            if allocated>cents(data['settlement_amount']):continue
            body.update({k:v for k,v in data.items() if k not in ['source','source_row']})
            c.execute('UPDATE records SET body=? WHERE id=?',(json.dumps(body),row['id']))
            updated+=1
    return updated

@app.get('/api/review')
def review():
    with db() as c:
        rs=records(c)
        blocked=[tuple(r) for r in c.execute('SELECT payment,receipt FROM match_blocks')]
        aliases=[tuple(r) for r in c.execute('SELECT payment_identity,receipt_identity FROM supplier_aliases')]
        suggestions=review_matches(rs,blocked,aliases)
        months={}
        for r in rs:
            if r['kind']!='payment':continue
            month=r.get('date','')[:7] or 'Date missing'
            m=months.setdefault(month,dict(payments=0,principal=0,fees=0,linked=0,remaining=0))
            m['payments']+=1
            for k,v in [('principal',r['amount']),('fees',r.get('fee',0)),('linked',r['allocated']),('remaining',r['remaining'])]:m[k]=round(m[k]+float(v or 0),2)
        return jsonify(suggestions=suggestions,months=months,
                       non_receipt_count=sum(r.get('collection')=='Non Receipt' for r in rs),
                       missing_dates=sum(not r.get('date') for r in rs),
                       missing_receipt_numbers=sum(r['kind']=='receipt' and r.get('collection')!='Non Receipt' and not r.get('receipt_no') for r in rs))

@app.post('/api/confirm-group')
def confirm_group():
    b=request.get_json()
    try:
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            rs=records(c)
            blocked=[tuple(r) for r in c.execute('SELECT payment,receipt FROM match_blocks')]
            aliases=[tuple(r) for r in c.execute('SELECT payment_identity,receipt_identity FROM supplier_aliases')]
            proposal=next((p for p in review_matches(rs,blocked,aliases) if p['receipt']==b.get('receipt') and sorted(p['payments'])==sorted(b.get('payments',[]))),None)
            if not proposal:raise ValueError('This suggestion has changed. Refresh the review queue.')
            if not proposal['can_confirm']:raise ValueError('Payment total exceeds receipt balance. Review allocations manually.')
            by_id={r['id']:r for r in rs}
            receipt=by_id[proposal['receipt']]
            for pid in proposal['payments']:
                payment=by_id[pid]
                value=cents(payment['remaining'])
                c.execute('INSERT INTO allocations VALUES(?,?,?) ON CONFLICT(payment,receipt) DO UPDATE SET cents=cents+excluded.cents',(pid,receipt['id'],value))
                c.execute('INSERT OR REPLACE INTO match_audit(payment,receipt,reason) VALUES(?,?,?)',(pid,receipt['id'],'Reviewed group: '+proposal['reason']+'; '+ '; '.join(proposal['issues'])))
                if b.get('remember_alias') is True:
                    c.execute('INSERT OR IGNORE INTO supplier_aliases(payment_identity,receipt_identity) VALUES(?,?)',(identity(payment),identity(receipt)))
        return jsonify(ok=True,remaining_difference=proposal['difference'],matching=auto_match())
    except (ValueError,KeyError) as e:return jsonify(error=str(e)),400

def import_workbook(path, label=None):
    label = label or Path(path).name
    w = openpyxl.load_workbook(path, read_only=True, data_only=True)
    count = 0
    with db() as c:
        scope=c.execute("SELECT value FROM workspace_settings WHERE key='scope_month'").fetchone()
        scope=scope['value'] if scope else None
        for s in w:
            if s.title in ['Suppliers', 'Suppliers List']:
                for row in s.iter_rows(min_row=2, values_only=True):
                    if len(row) >= 3 and isinstance(row[0], str) and row[0].strip():
                        c.execute('INSERT OR IGNORE INTO suppliers VALUES(?,?,?)', tuple(clean(x) for x in row[:3]))
                continue
            if s.title in ['COA', 'Copy of Accounts']:
                for row in s.iter_rows(values_only=True):
                    for i, v in enumerate(row[:-1]):
                        if re.fullmatch(r'[A-Z]+-\d{4}', clean(v)):
                            c.execute('INSERT OR IGNORE INTO accounts VALUES(?,?)', (clean(v), clean(row[i+1])))
                continue
            kind = 'receipt' if 'Voucher' in s.title or s.title == 'Non Receipt' else 'payment'
            if kind == 'payment' and s.title not in ['Material Expense', 'Cost of Sales Expense', 'Operating Expense', 'COS_Allowance']: continue
            headers = None
            for n, row in enumerate(s.iter_rows(values_only=True), 1):
                vals = [clean(x) for x in row]
                if headers is None:
                    if 'DATE' in vals and any(x in vals for x in ['PAYEE', 'PARTICULARS', 'PATRICULARS', 'EMPLOYEE NAME']): headers = vals
                    continue
                def get(*names):
                    for name in names:
                        if name in headers:
                            i = headers.index(name)
                            if i < len(vals): return vals[i]
                    return ''
                company = get('PAYEE', 'PARTICULARS', 'PATRICULARS', 'EMPLOYEE NAME')
                amount = get('AMOUNT', 'ALLOWANCE AMOUNT')
                if not company or not amount or company.upper() in ['TOTAL', 'GRAND TOTAL']: continue
                try: value = cents(amount)/100
                except ValueError: continue
                if value <= 0: continue
                if kind=='payment' and scope and not get('DATE').startswith(scope):continue
                b = dict.fromkeys(FIELDS, '')
                b.update(date=get('DATE'), company=company, reference=get('REF. NO', 'REFERENCE NO.', 'CHECK #'), method=get('MODE OF PAYMENT') or ('Cheque' if 'Check' in s.title else 'Cash'), tin=get('TIN'), address=get('ADDRESS'), amount=value, fee=cents(get('FEES'))/100, remarks=get('DESCRIPTION'), code=get('ACCOUNT CODE', 'CODE'), account=get('ACCOUNT TITLE', 'ACCOUNT'), project=get('PROJECT', 'SITE LOCATION'), receipt_no=get('RECEIPT NO.', 'S.I#'))
                b['collection'] = s.title
                if kind == 'receipt':
                    paid = get('AMOUNT PAID', 'CASH IN BANK')
                    b['settlement_amount'] = cents(paid)/100 if paid else value
                    b['withholding'] = cents(get('EXPANDED WITHHOLDING TAX'))/100
                b['receipt_month'] = scope or ('2026-09' if 'SEPTEMBER' in label.upper() else '')
                cur = c.execute('INSERT OR IGNORE INTO records VALUES(?,?,?,?,?)', (uuid.uuid4().hex, kind, json.dumps(b), label+' / '+s.title, n))
                count += cur.rowcount
    w.close()
    snapshot=audit(path) | {'file':label}
    enrich_workbook(snapshot)
    kind='quarterly' if any(s['name']=='Material Expense' for s in snapshot['sheets']) else 'monthly'
    app.register_source(path,kind,snapshot)
    return count

@app.get('/')
def index(): return render_template('materials_presence.html',deployment_notice='Hosted preview only. Local financial data is not included, and changes cannot be saved until persistent storage and access control are configured.' if HOSTED_PREVIEW else None)

@app.get('/assets/<path:name>')
def public_asset(name):return send_from_directory(ROOT/'static',name)

@app.get('/api/health')
def health():return jsonify(app='tally-local',workspace=str(ROOT.resolve()),pid=os.getpid(),version=RUNNING_VERSION)

@app.get('/api/state')
def state():
    with db() as c:
        scope=c.execute("SELECT value FROM workspace_settings WHERE key='scope_month'").fetchone()
        return jsonify(scope_month=scope['value'] if scope else None,records=records(c), suppliers=[dict(r) for r in c.execute('SELECT * FROM suppliers ORDER BY name')], accounts=[dict(r) for r in c.execute('SELECT * FROM accounts ORDER BY code')], allocations=[dict(r) for r in c.execute('SELECT a.*, m.reason FROM allocations a LEFT JOIN match_audit m ON a.payment=m.payment AND a.receipt=m.receipt')])

@app.post('/api/record')
def save():
    b = request.get_json()
    if b.get('kind') not in ['payment', 'receipt']: return jsonify(error='Choose payment or receipt'), 400
    if not str(b.get('company', '')).strip(): return jsonify(error='Company / payee is required'), 400
    try:
        b['amount'] = cents(b.get('amount'))/100
        b['fee'] = cents(b.get('fee'))/100
        if b['amount'] <= 0: raise ValueError('Amount must be greater than zero')
    except ValueError as e: return jsonify(error=str(e)), 400
    rid = b.get('id') or uuid.uuid4().hex
    with db() as c:
        old = c.execute('SELECT * FROM records WHERE id=?', (rid,)).fetchone()
        scope=c.execute("SELECT value FROM workspace_settings WHERE key='scope_month'").fetchone()
        workflow=c.execute("SELECT value FROM workspace_settings WHERE key='workflow'").fetchone()
        flexible_materials=workflow and workflow['value']=='materials_intake_receipt_presence' and b.get('collection') in ['Material Expense','Cash Voucher']
        if flexible_materials:
            try:
                if b['kind']=='payment' or b.get('date'):date.fromisoformat(str(b.get('date','')))
                if b['kind']=='receipt':date.fromisoformat(str(b.get('receipt_month',''))+'-01')
            except ValueError:return jsonify(error='Enter a valid payment date and receipt collection month'),400
            if b['kind']=='payment' and b.get('method') not in ['Online Payment','Bank','GCash','Cheque','Cash']:return jsonify(error='Choose a supported payment mode'),400
        if scope and not flexible_materials:
            month=scope['value']
            if b['kind']=='payment' and not str(b.get('date','')).startswith(month):return jsonify(error=f'This workspace contains {month} payments only. Enter a payment date in that month.'),400
            if b['kind']=='receipt' and b.get('receipt_month',month)!=month:return jsonify(error=f'This workspace contains the {month} receipt collection only.'),400
        if b.get('image'):
            doc=c.execute('SELECT record_id FROM documents WHERE image=?',(b['image'],)).fetchone()
            if doc and doc['record_id'] and doc['record_id']!=rid:return jsonify(error='This image is already attached to another entry'),409
        if old and old['kind'] != b['kind']: return jsonify(error='Record type cannot change'), 400
        allocated = c.execute('SELECT COALESCE(SUM(cents),0) FROM allocations WHERE payment=? OR receipt=?', (rid,rid)).fetchone()[0]
        if cents(b['amount']) < allocated: return jsonify(error='Amount is below the confirmed allocations'), 400
        if b.get('reference') and b['kind']=='payment':
            for r in records(c):
                if r['id'] != rid and r['kind']=='payment' and r.get('reference')==b['reference']: return jsonify(error='This payment reference already exists'), 409
        from intake import duplicates
        if not old:
            existing=duplicates(b,records(c))
            if existing:return jsonify(error='This document appears to belong to an existing entry. Attach the image to that entry instead.',duplicates=existing),409
        body = json.loads(old['body']) if old else {}
        body.update({k:v for k,v in b.items() if k not in ['id','kind','allocated','remaining','source','source_row','linked_receipt_numbers','settlement_amount','withholding']})
        for field in ['input_tax_goods','input_tax_services']:
            if field in b:
                try:body[field]=cents(b[field])/100
                except ValueError as e:return jsonify(error=str(e)),400
        if float(body.get('input_tax_goods',0))+float(body.get('input_tax_services',0))>b['amount']:return jsonify(error='Recorded input taxes exceed the gross amount'),400
        if 'input_tax_goods' in b or 'input_tax_services' in b:body['net_tax_base']=round(b['amount']-float(body.get('input_tax_goods',0))-float(body.get('input_tax_services',0)),2)
        if body.get('collection')=='COS_Allowance':body.setdefault('treatment','advance')
        treatment=body.get('treatment','business')
        if treatment not in ['business','personal','advance','no_receipt']:return jsonify(error='Choose a valid expense treatment'),400
        if b.get('withholding') is not None:
            try:body['withholding']=cents(b['withholding'])/100
            except ValueError as e:return jsonify(error=str(e)),400
        if body.get('withholding',0)>b['amount']:return jsonify(error='Withholding exceeds gross amount'),400
        if treatment=='no_receipt' and b['kind']=='receipt':body['collection']='Non Receipt'
        if treatment=='personal':body['personal']=True
        elif b.get('confirm_business') is True:body['personal']=False
        if b['kind']=='receipt':body.setdefault('receipt_month',str(b.get('date',''))[:7] or '2026-09')
        if b['kind']=='receipt' and 'receipt_no' in b:
            body['receipt_numbers']=[v.strip() for v in re.split(r'[,;]',str(b['receipt_no'])) if v.strip() and v.strip()!='-']
        if not old or 'settlement_amount' not in body:
            body['settlement_amount'] = b['amount']-float(body.get('withholding',0))
        elif old and cents(b['amount']) != cents(json.loads(old['body']).get('amount')):
            body['settlement_amount'] = max(0, b['amount'] - float(body.get('withholding',0)))
        if 'withholding' in b and (not old or cents(b['withholding'])!=cents(json.loads(old['body']).get('withholding'))):body['settlement_amount']=b['amount']-body['withholding']
        if b.get('net_override') not in (None,''):
            try:body['settlement_amount']=cents(b['net_override'])/100
            except ValueError as e:return jsonify(error=str(e)),400
        refunded=c.execute('SELECT COALESCE(SUM(cents),0) FROM returns WHERE payment=?',(rid,)).fetchone()[0]
        if cents(body['settlement_amount']) < allocated+refunded:
            return jsonify(error='Net amount is below the confirmed allocations'), 400
        if old: c.execute('UPDATE records SET body=? WHERE id=?', (json.dumps(body),rid))
        else: c.execute('INSERT INTO records VALUES(?,?,?,?,?)', (rid,b['kind'],json.dumps(body),'Manual / image review',None))
        if b.get('image'):c.execute('UPDATE documents SET record_id=? WHERE image=? AND (record_id IS NULL OR record_id=?)',(rid,b['image'],rid))
    linked=app.refresh_material_presence() if hasattr(app,'refresh_material_presence') else 0
    return jsonify(id=rid, matching=auto_match(),automatic_receipt_links=linked)

@app.post('/api/upload')
def upload():
    f = request.files.get('file')
    if not f:return jsonify(error='Choose a file'),400
    ext = Path(f.filename).suffix.lower()
    if ext not in ['.png','.jpg','.jpeg','.webp','.xlsx']: return jsonify(error='Use PNG, JPG, WebP or XLSX'),400
    name = uuid.uuid4().hex+ext
    path = DATA/'uploads'/name
    f.save(path)
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    if ext=='.xlsx':
        try:
            with db() as c:
                if c.execute('SELECT 1 FROM templates WHERE sha256=?',(digest,)).fetchone():return jsonify(imported=0,matching=auto_match(),note='Identical workbook already imported')
            imported = import_workbook(path, Path(f.filename).name)
            return jsonify(imported=imported, matching=auto_match())
        except Exception: return jsonify(error='Workbook could not be imported. Check that it uses the supported expense or voucher sheets.'),400
    from PIL import Image
    try:
        with Image.open(path) as img:img.verify()
    except Exception:return jsonify(error='This file is not a readable image'),400
    with db() as c:
        old=c.execute('SELECT * FROM documents WHERE sha256=?',(digest,)).fetchone()
        if old:return jsonify(image=old['image'],document_id=old['id'],record_id=old['record_id'],duplicate_image=True)
        did=uuid.uuid4().hex
        c.execute('INSERT INTO documents(id,image,sha256,original_name) VALUES(?,?,?,?)',(did,name,digest,Path(f.filename).name))
    return jsonify(image=name,document_id=did)

@app.get('/uploads/<name>')
def image(name): return send_from_directory(DATA/'uploads', name)

@app.post('/api/allocate')
def allocate():
    b=request.get_json()
    try:
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            rs={r['id']:r for r in records(c)}
            p,r=rs[b['payment']],rs[b['receipt']]
            if p['kind']!='payment' or r['kind']!='receipt': raise ValueError('Select a payment and a receipt')
            value=cents(b['amount'])
            if value<=0 or value>min(cents(p['remaining']),cents(r['remaining'])): raise ValueError('Allocation exceeds the remaining payment or receipt amount')
            c.execute('INSERT INTO allocations VALUES(?,?,?) ON CONFLICT(payment,receipt) DO UPDATE SET cents=cents+excluded.cents', (p['id'],r['id'],value))
            c.execute('DELETE FROM match_blocks WHERE payment=? AND receipt=?',(p['id'],r['id']))
            c.execute('INSERT OR REPLACE INTO match_audit(payment,receipt,reason) VALUES(?,?,?)',(p['id'],r['id'],'Manually confirmed allocation'))
        return jsonify(ok=True)
    except (ValueError, KeyError) as e: return jsonify(error=str(e)),400

@app.post('/api/unlink')
def unlink():
    b=request.get_json()
    with db() as c:
        c.execute('DELETE FROM allocations WHERE payment=? AND receipt=?',(b['payment'],b['receipt']))
        c.execute('DELETE FROM match_audit WHERE payment=? AND receipt=?',(b['payment'],b['receipt']))
        c.execute('INSERT OR IGNORE INTO match_blocks VALUES(?,?)',(b['payment'],b['receipt']))
    return jsonify(ok=True)

@app.get('/api/export')
def export():
    w=openpyxl.Workbook();w.remove(w.active)
    with db() as c:
        rs=records(c)
        for kind in ['payment','receipt']:
            s=w.create_sheet(kind.title()+'s')
            fields=FIELDS+['settlement_amount','withholding','linked_receipt_numbers','allocated','remaining','source','source_row']
            s.append(fields)
            for r in rs:
                if r['kind']==kind:
                    values=[r.get(k,'') for k in fields]
                    s.append(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in values])
            s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        s=w.create_sheet('Allocations');s.append(['payment_id','receipt_id','amount'])
        for a in c.execute('SELECT * FROM allocations'):s.append([a['payment'],a['receipt'],a['cents']/100])
        def safe(value):
            return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value
        s=w.create_sheet('Review Queue')
        s.append(['receipt_id','receipt_numbers','supplier','payment_ids','payment_total','net_receipt_balance','difference','reason','issues'])
        by_id={r['id']:r for r in rs}
        blocked=[tuple(r) for r in c.execute('SELECT payment,receipt FROM match_blocks')]
        aliases=[tuple(r) for r in c.execute('SELECT payment_identity,receipt_identity FROM supplier_aliases')]
        for p in review_matches(rs,blocked,aliases):
            r=by_id[p['receipt']]
            s.append([safe(v) for v in [r['id'],r.get('receipt_no',''),r['company'],', '.join(p['payments']),p['payment_total'],p['receipt_total'],p['difference'],p['reason'],'; '.join(p['issues'])]])
        for name,query in [('Supplier Master','SELECT name,tin,address FROM suppliers ORDER BY name'),('Chart of Accounts','SELECT code,title FROM accounts ORDER BY code')]:
            s=w.create_sheet(name)
            cursor=c.execute(query);s.append([d[0] for d in cursor.description])
            for row in cursor:s.append([safe(v) for v in row])
        s=w.create_sheet('Payment Monthly Summary')
        s.append(['month','expense_sheet','method','principal','fees','linked','remaining'])
        groups={}
        for r in rs:
            if r['kind']!='payment':continue
            key=(r.get('date','')[:7] or 'Date missing',r.get('collection','Manual'),r.get('method',''))
            totals=groups.setdefault(key,[0,0,0,0])
            for i,v in enumerate([r['amount'],r.get('fee',0),r['allocated'],r['remaining']]):totals[i]+=cents(v)
        for key,values in sorted(groups.items()):s.append([safe(v) for v in key]+[v/100 for v in values])
    buf=io.BytesIO();w.save(buf);buf.seek(0)
    return send_file(buf,download_name='reconciled-expenses.xlsx',as_attachment=True)

from system import install
install(app,lambda:DATA,db,records,cents,auto_match)
from modules import install as install_modules
install_modules(app,db,records,cents)
from presence import install as install_presence
install_presence(app,db,records)
from materials import install as install_materials
install_materials(app,lambda:DATA,db,records,cents)

if __name__=='__main__':
    (DATA/'server.pid').write_text(str(os.getpid()))
    app.run(host='127.0.0.1',port=5050,debug=False)
