import io
import json
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
import openpyxl
import app as module


class ReviewApiTests(unittest.TestCase):
    def setUp(self):
        self.original=module.DATA
        self.temp=tempfile.TemporaryDirectory()
        self.data=Path(self.temp.name)
        with module.db() as source, closing(sqlite3.connect(self.data/'receipts.sqlite3')) as target:
            source.backup(target)
        module.DATA=self.data
        with module.db() as c:
            c.execute("DELETE FROM workspace_settings WHERE key='workflow'")
            for table in ['presence_links','presence_dismissals','allocations','match_audit','match_blocks','supplier_aliases','documents','returns','records']:
                c.execute('DELETE FROM '+table)
            for rid,kind,amount,name,si in [('p1','payment',15120,'Payment Alias','35522'),('p2','payment',1120,'Payment Alias','35522'),('r','receipt',16240,'Supplier Legal Name','35522')]:
                body=dict(company=name,amount=amount,fee=0,receipt_no=si,settlement_amount=amount)
                c.execute('INSERT INTO records VALUES(?,?,?,?,?)',(rid,kind,json.dumps(body),'Fixture',None))
        self.client=module.app.test_client()

    def tearDown(self):
        module.DATA=self.original
        self.temp.cleanup()

    def test_confirm_alias_undo_and_export(self):
        self.assertEqual(len(self.client.get('/api/review').get_json()['suggestions']),1)
        response=self.client.post('/api/confirm-group',json={'receipt':'r','payments':['p1','p2'],'remember_alias':True})
        self.assertEqual(response.status_code,200)
        with module.db() as c:
            self.assertEqual(c.execute('SELECT COUNT(*) FROM supplier_aliases').fetchone()[0],1)
            self.assertEqual(sum(r['remaining'] for r in module.records(c)),0)
        self.assertEqual(self.client.post('/api/auto-match',json={}).get_json()['links'],0)
        self.assertEqual(self.client.post('/api/unlink',json={'payment':'p1','receipt':'r'}).status_code,200)
        self.assertEqual(self.client.post('/api/auto-match',json={}).get_json()['links'],0)
        response=self.client.get('/api/export')
        self.assertEqual(response.status_code,200)
        workbook=openpyxl.load_workbook(io.BytesIO(response.data),read_only=True)
        self.assertTrue({'Review Queue','Supplier Master','Chart of Accounts','Payment Monthly Summary'}.issubset(workbook.sheetnames))
        workbook.close()

    def test_stale_or_oversized_group_rejected(self):
        response=self.client.post('/api/confirm-group',json={'receipt':'r','payments':['p1']})
        self.assertEqual(response.status_code,400)
        with module.db() as c:
            r=json.loads(c.execute("SELECT body FROM records WHERE id='r'").fetchone()[0]);r['settlement_amount']=16000
            c.execute("UPDATE records SET body=? WHERE id='r'",(json.dumps(r),))
        self.assertEqual(self.client.post('/api/confirm-group',json={'receipt':'r','payments':['p1','p2']}).status_code,400)
        with module.db() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM allocations').fetchone()[0],0)

    def test_all_fourteen_modules_and_shared_masters(self):
        definitions=self.client.get('/api/modules').get_json()['modules']
        self.assertEqual(len(definitions),14)
        for definition in definitions:
            response=self.client.get('/api/modules/'+definition['id'])
            self.assertEqual(response.status_code,200,definition['id'])
        response=self.client.post('/api/master/supplier',json={'name':'New Supplier Fixture','tin':'123-456-789','address':'Fixture Street'})
        self.assertEqual(response.status_code,200)
        for slug in ['supplier-list','suppliers']:
            self.assertTrue(any(s['name']=='New Supplier Fixture' for s in self.client.get('/api/modules/'+slug).get_json()['rows']))
        self.assertEqual(self.client.post('/api/master/account',json={'code':'CS-9999','title':'Fixture Expense'}).status_code,200)
        self.assertEqual(self.client.post('/api/master/account',json={'code':'CS-9999','title':'Duplicate'}).status_code,409)

    def test_photo_intake_duplicate_and_attach(self):
        from PIL import Image
        image=io.BytesIO();Image.new('RGB',(3,3),'red').save(image,format='PNG');image=image.getvalue()
        (module.DATA/'uploads').mkdir(exist_ok=True)
        first=self.client.post('/api/upload',data={'file':(io.BytesIO(image),'proof.png')},content_type='multipart/form-data').get_json()
        again=self.client.post('/api/upload',data={'file':(io.BytesIO(image),'same.png')},content_type='multipart/form-data').get_json()
        self.assertTrue(again['duplicate_image']);self.assertEqual(first['image'],again['image'])
        draft=self.client.post('/api/intake/parse',json={'mode':'payment','text':'GCash\nSent to:\nPhoto Supplier\nAmount\nPHP 50.00\nReference No. ABC123456\nSep 10, 2026','month':'2026-09'}).get_json()['draft']
        self.assertEqual(draft['amount'],'50.00');self.assertEqual(draft['reference'],'ABC123456')
        draft['image']=first['image'];draft['company']='Photo Supplier'
        saved=self.client.post('/api/record',json=draft);self.assertEqual(saved.status_code,200,saved.data)
        rid=saved.get_json()['id']
        self.assertEqual(len(self.client.get('/api/documents/'+rid).get_json()['documents']),1)
        self.assertEqual(self.client.post('/api/record',json=draft).status_code,409)
        self.assertEqual(self.client.post('/api/intake/attach',json={'record_id':'p1','image':first['image']}).status_code,409)

    def test_advance_return_and_summary_no_double_count(self):
        b=dict(kind='payment',company='Advance Fixture',date='2026-09-10',amount=1000,fee=10,treatment='advance',collection='COS_Allowance')
        response=self.client.post('/api/record',json=b);self.assertEqual(response.status_code,200)
        rid=response.get_json()['id']
        response=self.client.post('/api/returns',json={'payment':rid,'amount':400,'date':'2026-09-12','reason':'Unused advance returned'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.client.post('/api/returns',json={'payment':rid,'amount':601,'date':'2026-09-12','reason':'Too much'}).status_code,400)
        rows=self.client.get('/api/modules/allowances').get_json()['rows']
        self.assertEqual(next(r for r in rows if r['id']==rid)['remaining'],600)
        self.client.post('/api/controls',json={'date':'2026-09-10','kind':'cash','amount':99999,'note':'Aggregate control'})
        summaries=self.client.get('/api/modules/daily').get_json()['rows']
        summary=next(r for r in summaries if r['period']=='2026-09-10')
        self.assertEqual(summary['total'],1010)

    def test_scoped_workspace_rejects_other_months(self):
        with module.db() as c:c.execute("INSERT OR REPLACE INTO workspace_settings VALUES('scope_month','2026-09')")
        self.assertEqual(self.client.post('/api/record',json={'kind':'payment','company':'Outside','amount':100,'date':'2026-08-31'}).status_code,400)
        self.assertEqual(self.client.post('/api/controls',json={'kind':'cash','date':'2026-08-31','amount':100}).status_code,400)
        self.assertEqual(self.client.get('/api/export/complete?month=2026-08').status_code,400)
        self.assertEqual(self.client.get('/api/state').get_json()['scope_month'],'2026-09')

    def test_material_presence_decision_does_not_allocate_money(self):
        with module.db() as c:
            for rid in ['p1','r']:
                body=json.loads(c.execute('SELECT body FROM records WHERE id=?',(rid,)).fetchone()[0])
                body.update(collection='Material Expense' if rid=='p1' else 'September Cash Voucher',method='Online Payment' if rid=='p1' else 'Cash',code='CS-0001',account='Materials',date='2026-09-01' if rid=='p1' else '2026-07-01')
                c.execute('UPDATE records SET body=? WHERE id=?',(json.dumps(body),rid))
        response=self.client.get('/api/materials-presence');self.assertEqual(response.status_code,200)
        self.assertEqual(len(response.get_json()['rows']),1)
        response=self.client.post('/api/materials-presence/decision',json={'payment':'p1','receipt':'r','action':'confirm'})
        self.assertEqual(response.status_code,200)
        with module.db() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM allocations').fetchone()[0],0)
        self.assertEqual(self.client.get('/api/materials-presence').get_json()['rows'][0]['status'],'found')
        self.assertEqual(self.client.post('/api/materials-presence/decision',json={'payment':'p1','receipt':'r','action':'dismiss'}).status_code,200)
        self.assertEqual(self.client.get('/api/materials-presence').get_json()['rows'][0]['status'],'not_found')
        response=self.client.get('/api/materials-presence/export');self.assertEqual(response.status_code,200)
        w=openpyxl.load_workbook(io.BytesIO(response.data),read_only=True)
        self.assertIn('No candidate found',w.sheetnames);w.close()

    def test_all_material_payment_modes_and_future_dates(self):
        with module.db() as c:c.execute("INSERT OR REPLACE INTO workspace_settings VALUES('workflow','materials_intake_receipt_presence')")
        for i,method in enumerate(['Online Payment','Bank','GCash','Cheque','Cash']):
            r=self.client.post('/api/record',json=dict(kind='payment',collection='Material Expense',date='2026-10-05',company='Mode Fixture '+method,method=method,amount=100+i,fee=5,reference='MODE-TEST-'+str(i),code='CS-0001',account='Materials'))
            self.assertEqual(r.status_code,200,r.data)
        rows=self.client.get('/api/materials/entries?month=2026-10').get_json()['rows']
        self.assertEqual(len(rows),5)
        self.assertEqual({r['method'] for r in rows},{'Online Payment','Bank','GCash','Cheque','Cash'})
        response=self.client.get('/api/materials/export?month=2026-10')
        w=openpyxl.load_workbook(io.BytesIO(response.data),read_only=True,data_only=True)
        self.assertEqual(w['Material Expense'].max_column,22)
        self.assertEqual(w['Material Expense']['N2'].value,105)
        w.close()

    def test_reference_workbook_drives_image_text_fields(self):
        w=openpyxl.Workbook();s=w.active;s.title='Suppliers List'
        s.append(['Supplier','Tin','Address']);s.append(['OCR Supplier Fixture','123-456-789','Fixture Street'])
        a=w.create_sheet('Copy of Accounts');a.append(['CS-0001','Materials'])
        buf=io.BytesIO();w.save(buf);w.close()
        r=self.client.post('/api/materials/import',data={'kind':'references','month':'2026-10','file':(io.BytesIO(buf.getvalue()),'references.xlsx')},content_type='multipart/form-data')
        self.assertEqual(r.status_code,200,r.data)
        self.assertEqual(r.get_json()['supplier_rows'],1)
        r=self.client.post('/api/materials/parse',json={'mode':'payment','month':'2026-10','text':'Sent to:\nOCR Supplier Fixture\nAmount\nPHP 1,000.00\nCheque No. 888888\nOct 5, 2026'})
        self.assertEqual(r.status_code,200)
        draft=r.get_json()['draft']
        self.assertEqual(draft['company'],'OCR Supplier Fixture')
        self.assertEqual(draft['tin'],'123-456-789')
        self.assertEqual(draft['address'],'Fixture Street')
        self.assertEqual(draft['code'],'CS-0001')
        self.assertEqual(draft['method'],'Cheque')
        self.assertEqual(draft['reference'],'888888')

    def test_next_month_voucher_import_matches_older_payment_without_totals(self):
        with module.db() as c:c.execute("INSERT OR REPLACE INTO workspace_settings VALUES('workflow','materials_intake_receipt_presence')")
        p=self.client.post('/api/record',json=dict(kind='payment',collection='Material Expense',date='2026-09-20',company='Voucher Upload Fixture',tin='123-456-789',method='GCash',amount=1000,fee=15,reference='GCASH-VOUCHER-TEST',receipt_no='INV-TEST-001'))
        self.assertEqual(p.status_code,200,p.data);pid=p.get_json()['id']
        w=openpyxl.Workbook();s=w.active;s.title='October Cash Voucher'
        headers={'A':'DATE','D':'PARTICULARS','M':'S.I#','N':'TIN','O':'ADDRESS','P':'CODE','Q':'ACCOUNT','R':'TAG','S':'AMOUNT','T':'INPUT TAX - GOODS','U':'INPUT TAX - SERVICES','V':'NET AMOUNT','W':'EXPANDED WITHHOLDING TAX','X':'AMOUNT PAID'}
        for col,value in headers.items():s[col+'1']=value
        from datetime import datetime
        values={'A':datetime(2026,7,10),'D':'Voucher Upload Fixture','M':'INV-TEST-001','N':'123-456-789','P':'CS-0001','Q':'Materials','S':1000,'W':0,'X':1000}
        for col,value in values.items():s[col+'2']=value
        buf=io.BytesIO();w.save(buf);w.close();content=buf.getvalue()
        def upload():return self.client.post('/api/materials/import',data={'kind':'vouchers','month':'2026-10','file':(io.BytesIO(content),'October.xlsx')},content_type='multipart/form-data')
        response=upload();self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.get_json()['imported'],1)
        self.assertEqual(response.get_json()['automatic_receipt_links'],1)
        data=self.client.get('/api/materials-presence?month=2026-10&method=GCash').get_json()
        self.assertEqual(data['rows'][0]['status'],'found')
        self.assertTrue(data['rows'][0]['candidates'][0]['automatic'])
        self.assertEqual(data['vouchers'][0]['date'],'2026-07-10')
        self.assertEqual(upload().get_json()['imported'],0)
        entries=self.client.get('/api/materials/entries').get_json()['rows']
        self.assertEqual(next(r for r in entries if r['id']==pid)['matched_receipt_numbers'],'INV-TEST-001')
        with module.db() as c:self.assertEqual(c.execute('SELECT SUM(cents) FROM allocations WHERE payment=?',(pid,)).fetchone()[0],100000)
        # A recurring monthly filename/sheet must not collide with prior months.
        november=self.client.post('/api/materials/import',data={'kind':'vouchers','month':'2026-11','file':(io.BytesIO(content),'October.xlsx')},content_type='multipart/form-data')
        self.assertEqual(november.status_code,200,november.data)
        self.assertEqual(november.get_json()['imported'],1)
        self.assertEqual(len(self.client.get('/api/materials-presence?month=2026-11').get_json()['vouchers']),1)




    def test_workbook_entry_registers_and_multiple_receipts_per_day(self):
        with module.db() as c:
            c.execute("INSERT OR REPLACE INTO workspace_settings VALUES('workflow','materials_intake_receipt_presence')")
        self.assertEqual(self.client.get('/workbook').status_code,200)
        for collection,slug,method,count in [('Cash Voucher','cash-vouchers','Cash',3),('Check Voucher','cheque-vouchers','Cheque',1),('Non Receipt','non-receipt','GCash',1)]:
            for n in range(count):
                body=dict(kind='receipt',collection=collection,receipt_month='2026-10',date='2026-07-02',company='Example '+collection,amount=100+n,method=method,code='CS-0001',receipt_no='' if collection=='Non Receipt' else 'SI-'+str(n),treatment='no_receipt' if collection=='Non Receipt' else 'business')
                response=self.client.post('/api/record',json=body)
                self.assertEqual(response.status_code,200,response.get_json())
            rows=self.client.get('/api/modules/'+slug+'?month=2026-10').get_json()['rows']
            self.assertEqual(len(rows),count)
            self.assertTrue(all(r['date']=='2026-07-02' for r in rows))
        for collection in ['Cost of Sales Expense','Operating Expense','COS_Allowance']:
            response=self.client.post('/api/record',json=dict(kind='payment',collection=collection,date='2026-10-24',company='Example employee',amount=200,method='Cash'))
            self.assertEqual(response.status_code,200,response.get_json())
        response=self.client.post('/api/controls',json=dict(kind='gcash',date='2026-10-24',amount=600,note='Daily control'))
        self.assertEqual(response.status_code,200,response.get_json())
        self.assertEqual(len(self.client.get('/api/modules/cash-controls?month=2026-10').get_json()['rows']),1)
        self.assertEqual(len(self.client.get('/api/modules/cash-vouchers?month=2026-11').get_json()['rows']),0)

    def test_company_personal_payment_split(self):
        with module.db() as c:
            c.execute("INSERT OR REPLACE INTO workspace_settings VALUES('workflow','materials_intake_receipt_presence')")
        body=dict(kind='payment',collection='Material Expense',date='2026-10-23',company='Example supplier',amount=25133.48,transfer_amount=25133.48,personal_amount=1100,personal_note='Personal beneficiary',method='GCash',reference='SPLIT-TEST-001',code='CS-0001')
        response=self.client.post('/api/record',json=body)
        self.assertEqual(response.status_code,200,response.get_json())
        rid=response.get_json()['id']
        entry=next(r for r in self.client.get('/api/materials/entries').get_json()['rows'] if r['id']==rid)
        for field in ['amount','settlement_amount','remaining']:self.assertEqual(entry[field],24033.48)
        self.assertEqual(entry['transfer_amount'],25133.48)
        self.assertEqual(entry['personal_amount'],1100)
        october=next(r for r in self.client.get('/api/modules/monthly').get_json()['rows'] if r['period']=='2026-10')
        self.assertEqual(october['total'],24033.48)
        self.assertEqual(october['personal'],1100)
        body.update(id=rid,personal_amount=26000)
        self.assertEqual(self.client.post('/api/record',json=body).status_code,400)
        body.update(personal_amount=1100,amount=24033.48)
        self.assertEqual(self.client.post('/api/record',json=body).status_code,200)

    def test_partial_payments_find_si_without_auto_confirming_amounts(self):
        with module.db() as c:
            c.execute("INSERT OR REPLACE INTO workspace_settings VALUES('workflow','materials_intake_receipt_presence')")
            for rid,kind,amount in [('exact1','payment',100),('exact2','payment',200),('exactr','receipt',300)]:
                body=dict(company='Example supplier',amount=amount,receipt_no='SI-987',collection='Material Expense' if kind=='payment' else 'Cash Voucher',code='CS-0001',date='2026-09-01',receipt_month='2026-09')
                c.execute('INSERT INTO records VALUES(?,?,?,?,?)',(rid,kind,json.dumps(body),'Example',None))
        module.app.refresh_material_presence()
        rows=self.client.get('/api/materials-presence?month=2026-09').get_json()['rows']
        for row in rows:
            if row['payment']['id'] in ['exact1','exact2']:
                self.assertEqual(row['status'],'found')
                self.assertFalse(next(c for c in row['candidates'] if c['receipt']['id']=='exactr').get('automatic',False))
        with module.db() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM allocations').fetchone()[0],0)

    def test_auto_link_identity_and_amount_required(self):
        with module.db() as c:
            for rid,kind,reference,amount,tin in [('refp','payment','PC-EXAMPLE-001',100,'123-456-789'),('refr','receipt','PC-EXAMPLE-001',100,'123-456-789')]:
                body=dict(company='Example reference supplier',amount=amount,reference=reference,tin=tin,collection='Material Expense' if kind=='payment' else 'Cash Voucher',code='CS-0001',receipt_month='2026-09')
                c.execute('INSERT INTO records VALUES(?,?,?,?,?)',(rid,kind,json.dumps(body),'Example',None))
        self.assertEqual(module.app.refresh_material_presence(),1)
        for field,value in [('amount',101),('tin','987-654-321')]:
            with module.db() as c:
                original=json.loads(c.execute("SELECT body FROM records WHERE id='refr'").fetchone()[0]);changed=dict(original);changed[field]=value
                c.execute("UPDATE records SET body=? WHERE id='refr'",(json.dumps(changed),))
            self.assertEqual(module.app.refresh_material_presence(),0)
            with module.db() as c:c.execute("UPDATE records SET body=? WHERE id='refr'",(json.dumps(original),))

    def test_auto_allocation_uses_strict_rule_in_current_workflow(self):
        with module.db() as c:
            c.execute("INSERT OR REPLACE INTO workspace_settings VALUES('workflow','materials_intake_receipt_presence')")
            for rid,kind in [('allocp','payment'),('allocr','receipt')]:
                body=dict(company='Example allocation supplier',amount=100,settlement_amount=100,reference='PC-ALLOC-001',collection='Material Expense' if kind=='payment' else 'Cash Voucher',code='CS-0001',receipt_month='2026-09')
                c.execute('INSERT INTO records VALUES(?,?,?,?,?)',(rid,kind,json.dumps(body),'Example',None))
        self.assertEqual(self.client.post('/api/auto-match',json={}).get_json()['links'],3)
        self.assertEqual(self.client.post('/api/auto-match',json={}).get_json()['links'],0)
        with module.db() as c:
            self.assertEqual(c.execute("SELECT cents FROM allocations WHERE payment='allocp' AND receipt='allocr'").fetchone()[0],10000)

    def test_unique_supplier_amount_fills_si_and_duplicate_stays_review(self):
        with module.db() as c:
            c.execute("INSERT OR REPLACE INTO workspace_settings VALUES('workflow','materials_intake_receipt_presence')")
            for rid,kind,si in [('uniquep','payment',''),('uniquer','receipt','SI-UNIQUE-001')]:
                body=dict(company='Example unique supplier',amount=123.45,settlement_amount=123.45,receipt_no=si,collection='Material Expense' if kind=='payment' else 'Cash Voucher',code='CS-0001',receipt_month='2026-09')
                c.execute('INSERT INTO records VALUES(?,?,?,?,?)',(rid,kind,json.dumps(body),'Example',None))
        self.assertEqual(module.app.refresh_material_presence(),1)
        self.assertEqual(module.auto_match()['links'],3)
        entry=next(r for r in self.client.get('/api/materials/entries').get_json()['rows'] if r['id']=='uniquep')
        self.assertEqual(entry['matched_receipt_numbers'],'SI-UNIQUE-001')
        self.assertEqual(entry['remaining'],0)
        with module.db() as c:
            for rid,kind,si in [('ambp','payment',''),('ambr1','receipt','SI-A'),('ambr2','receipt','SI-B')]:
                body=dict(company='Example ambiguous supplier',amount=200,receipt_no=si,collection='Material Expense' if kind=='payment' else 'Cash Voucher',code='CS-0001',receipt_month='2026-09')
                c.execute('INSERT INTO records VALUES(?,?,?,?,?)',(rid,kind,json.dumps(body),'Example',None))
        self.assertEqual(module.auto_match()['links'],0)
        module.app.refresh_material_presence()
        row=next(r for r in self.client.get('/api/materials-presence').get_json()['rows'] if r['payment']['id']=='ambp')
        self.assertEqual(row['status'],'review')

    def test_proprietor_company_mapping_is_remembered(self):
        response=self.client.post('/api/master/supplier',json=dict(name='Example Owner - Prop.',business_name='Example Hardware',proprietor='Example Owner',tin='111-222-333',address='Example address'))
        self.assertEqual(response.status_code,200,response.get_json())
        enriched=self.client.post('/api/materials/enrich',json=dict(company='Example Hardware')).get_json()['fields']
        self.assertEqual(enriched['company'],'Example Owner - Prop.')
        self.assertEqual(enriched['business_name'],'Example Hardware')
        self.assertEqual(enriched['proprietor'],'Example Owner')
        with module.db() as c:
            for rid,kind,company in [('ownerp','payment','Example Hardware'),('ownerr','receipt','Example Owner - Prop.')]:
                body=dict(company=company,amount=333,collection='Material Expense' if kind=='payment' else 'Cash Voucher',receipt_no='' if kind=='payment' else 'SI-OWNER-1',code='CS-0001',receipt_month='2026-09')
                c.execute('INSERT INTO records VALUES(?,?,?,?,?)',(rid,kind,json.dumps(body),'Example',None))
        self.assertEqual(module.auto_match()['links'],1)
        saved=self.client.get('/api/modules/suppliers?q=example%20hardware').get_json()['rows']
        self.assertEqual(saved[0]['proprietor'],'Example Owner')
        parsed=self.client.post('/api/materials/parse',json=dict(text='Example Hardware\nAmount 333.00',mode='payment',month='2026-09')).get_json()['draft']
        self.assertEqual(parsed['company'],'Example Owner - Prop.')

    def test_record_save_learns_explicit_proprietor_and_links(self):
        with module.db() as c:
            c.execute("INSERT OR REPLACE INTO workspace_settings VALUES('workflow','materials_intake_receipt_presence')")
            body=dict(company='Example New Hardware',amount=444,collection='Material Expense',date='2026-09-01',method='Bank')
            c.execute('INSERT INTO records VALUES(?,?,?,?,?)',('learnp','payment',json.dumps(body),'Example',None))
        response=self.client.post('/api/record',json=dict(kind='receipt',company='Example New Owner - Prop.',business_name='Example New Hardware',amount=444,collection='Cash Voucher',receipt_month='2026-09',receipt_no='SI-LEARN-1',code='CS-0001'))
        self.assertEqual(response.status_code,200,response.get_json())
        self.assertEqual(response.get_json()['matching']['links'],3)
        with module.db() as c:
            mapping=c.execute("SELECT * FROM supplier_profiles WHERE name='Example New Owner - Prop.'").fetchone()
            self.assertEqual(mapping['business_name'],'Example New Hardware')
        response=self.client.post('/api/materials/enrich',json={'company':'Example New Hardware'})
        self.assertEqual(response.get_json()['fields']['company'],'Example New Owner - Prop.')

    def test_proprietor_company_learned_from_unique_taxpayer(self):
        from supplier_profiles import learn_from_supplier_tins
        with module.db() as c:
            c.executemany('INSERT INTO suppliers VALUES(?,?,?)',[('Example Tax Owner - Prop.','987-111-222','Address'),('Example Tax Hardware','987-111-222-000','Address')])
            self.assertGreaterEqual(learn_from_supplier_tins(c),1)
            row=c.execute("SELECT business_name FROM supplier_profiles WHERE name='Example Tax Owner - Prop.'").fetchone()
            self.assertEqual(row['business_name'],'Example Tax Hardware')

    def test_existing_confirmed_si_teaches_payee_alias_and_allocates_other_receipts(self):
        with module.db() as c:
            for rid,kind,name,amount,si in [('historyp','payment','Example Electrical',100,'OLD-101'),('historyr','receipt','Example Owner - Prop.',100,'OLD-101'),('newp','payment','Example Electrical',200,''),('newr','receipt','Example Owner - Prop.',200,'NEW-202')]:
                body=dict(company=name,amount=amount,collection='Material Expense' if kind=='payment' else 'Cash Voucher',receipt_no=si,code='CS-0001',receipt_month='2026-09')
                c.execute('INSERT INTO records VALUES(?,?,?,?,?)',(rid,kind,json.dumps(body),'Example',None))
            c.execute('INSERT INTO allocations VALUES(?,?,?)',('historyp','historyr',10000))
        result=module.auto_match()
        with module.db() as c:
            self.assertEqual(c.execute("SELECT cents FROM allocations WHERE payment='newp' AND receipt='newr'").fetchone()[0],20000)
            self.assertEqual(c.execute("SELECT legal_name FROM supplier_name_links WHERE alias='exampleelectrical'").fetchone()[0],'Example Owner - Prop.')
        closed=self.client.get('/api/materials/entries?status=closed').get_json()['rows']
        row=next(r for r in closed if r['id']=='newp')
        self.assertEqual(row['matched_receipt_numbers'],'NEW-202')
        self.assertFalse(any(r['id']=='newp' for r in self.client.get('/api/materials/entries?status=open').get_json()['rows']))

    def test_lightweight_change_revision_detects_saved_changes(self):
        before=self.client.get('/api/changes')
        self.assertEqual(before.status_code,200)
        self.assertEqual(before.headers['Cache-Control'],'no-store')
        with module.db() as c:c.execute('INSERT OR REPLACE INTO accounts VALUES(?,?)',('ZZ-9999','Revision test'))
        after=self.client.get('/api/changes').get_json()
        self.assertNotEqual(before.get_json()['revision'],after['revision'])
        self.assertEqual(set(after),{'revision'})

    def test_rounding_difference_remains_visible_after_automatic_allocation(self):
        with module.db() as c:
            for rid,kind,amount in [('centp','payment',2983.13),('centr','receipt',2983.12)]:
                body=dict(company='Example cent supplier',amount=amount,receipt_no='SI-CENT-01',tin='333-444-555',collection='Material Expense' if kind=='payment' else 'Cash Voucher',code='CS-0001',receipt_month='2026-09')
                c.execute('INSERT INTO records VALUES(?,?,?,?,?)',(rid,kind,json.dumps(body),'Example',None))
        module.auto_match()
        rows=self.client.get('/api/materials/entries?status=linked').get_json()['rows']
        row=next(r for r in rows if r['id']=='centp')
        self.assertEqual(row['amount'],2983.13)
        self.assertEqual(row['allocated'],2983.12)
        self.assertEqual(row['remaining'],.01)
        self.assertEqual(row['matched_receipt_numbers'],'SI-CENT-01')
        self.assertEqual(module.auto_match()['links'],0)
