import io
import tempfile
import unittest
from pathlib import Path
import openpyxl
from modules import DEFINITIONS
from reports import build_workbook


class WholeWorkbookTests(unittest.TestCase):
    def test_all_source_tabs_totals_and_receipt_numbers(self):
        records=[dict(id='p',kind='payment',date='2026-09-10',company='Supplier',amount=1110,fee=0,allocated=1110,remaining=0,collection='Material Expense',method='Bank',code='CS-0001',account='Materials',linked_receipt_numbers='551, 7985'),
                 dict(id='r',kind='receipt',date='2026-08-10',receipt_month='2026-09',company='Supplier',amount=1120,settlement_amount=1110,withholding=10,input_tax_goods=120,input_tax_services=0,receipt_no='551, 7985',receipt_numbers=['551','7985'],collection='September Cash Voucher',method='Cash')]
        with tempfile.TemporaryDirectory() as td:
            for kind,count in [('quarterly',9),('monthly',5)]:
                original=openpyxl.Workbook();original.remove(original.active);snapshot={'sheets':[]}
                for name,workbook,_ in DEFINITIONS.values():
                    if workbook!=kind:continue
                    original.create_sheet(name)['A1']='Source header'
                    snapshot['sheets'].append(dict(name=name,rows=[dict(row=1,values={'A':'Source header'})]))
                path=Path(td)/(kind+'.xlsx');original.save(path);original.close()
                content,manifest=build_workbook(path,snapshot,records,[],[],[dict(kind='cash',date='2026-09-10',amount=99999)],'2026-09',kind=='quarterly')
                result=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=True)
                self.assertEqual(len(result.sheetnames),count)
                if kind=='quarterly':
                    self.assertEqual(result['Material Expense']['N4'].value,1110)
                    self.assertEqual(result['Material Expense']['C4'].value,'551, 7985')
                    self.assertEqual(result['Expense Summary']['I2'].value,1110)
                    self.assertEqual(manifest['payment_principal'],1110)
                else:
                    self.assertEqual(result['September Cash Voucher']['M14'].value,'551, 7985')
                    self.assertEqual(result['September Cash Voucher']['S14'].value,1120)
                    self.assertEqual(result['September Cash Voucher']['X14'].value,1110)
                    self.assertEqual(manifest['receipt_gross'],1120)
                result.close()


if __name__=='__main__':unittest.main()
